import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import timedelta
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.models import User
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import render, get_object_or_404
from django.utils import timezone
from ..models import RegistroAuditoria


def _es_superusuario(user):
    return user.is_authenticated and user.is_superuser


def _aplicar_filtros(request, registros):
    """Aplica los filtros de la URL al queryset de registros.
    Reusable desde la vista de lista y desde el export a Excel."""
    usuario_id = request.GET.get('usuario', '').strip()
    accion = request.GET.get('accion', '').strip()
    modelo = request.GET.get('modelo', '').strip()
    fecha_desde = request.GET.get('fecha_desde', '').strip()
    fecha_hasta = request.GET.get('fecha_hasta', '').strip()
    buscar = request.GET.get('buscar', '').strip()

    if usuario_id:
        registros = registros.filter(usuario_id=usuario_id)
    if accion:
        registros = registros.filter(accion=accion)
    if modelo:
        registros = registros.filter(modelo=modelo)
    if fecha_desde:
        registros = registros.filter(fecha__date__gte=fecha_desde)
    if fecha_hasta:
        registros = registros.filter(fecha__date__lte=fecha_hasta)
    if buscar:
        registros = registros.filter(
            Q(objeto_repr__icontains=buscar) |
            Q(ip__icontains=buscar) |
            Q(user_agent__icontains=buscar)
        )

    filtros = {
        'usuario_id': usuario_id,
        'accion': accion,
        'modelo': modelo,
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'buscar': buscar,
    }
    return registros, filtros


@login_required(login_url='login')
@user_passes_test(_es_superusuario, login_url='dashboard')
def auditoria_list(request):
    """Listado de registros de auditoría, solo para superusuarios."""
    registros = RegistroAuditoria.objects.select_related('usuario').all().order_by('-fecha')
    registros, filtros = _aplicar_filtros(request, registros)

    # Contadores por acción (sobre el queryset filtrado)
    contadores = {
        'CREAR': registros.filter(accion='CREAR').count(),
        'EDITAR': registros.filter(accion='EDITAR').count(),
        'ELIMINAR': registros.filter(accion='ELIMINAR').count(),
        'LOGIN': registros.filter(accion='LOGIN').count(),
        'LOGOUT': registros.filter(accion='LOGOUT').count(),
    }

    # Paginación
    paginator = Paginator(registros, 50)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    usuarios = User.objects.filter(registros_auditoria__isnull=False).distinct().order_by('username')
    acciones = RegistroAuditoria.ACCION_CHOICES
    modelos = (
        RegistroAuditoria.objects
        .values_list('modelo', flat=True)
        .distinct()
        .order_by('modelo')
    )

    hoy = timezone.localdate()

    context = {
        'page_obj': page_obj,
        'total_registros': registros.count(),
        'contadores': contadores,
        'usuarios': usuarios,
        'acciones': acciones,
        'modelos': modelos,
        'hoy': hoy.isoformat(),
        'primer_dia_mes': hoy.replace(day=1).isoformat(),
        'hace_7_dias': (hoy - timedelta(days=7)).isoformat(),
        **filtros,
    }
    return render(request, 'facturacion/auditoria/auditoria_list.html', context)


@login_required(login_url='login')
@user_passes_test(_es_superusuario, login_url='dashboard')
def auditoria_detalle(request, pk):
    """Detalle de un registro de auditoría (los cambios)."""
    registro = get_object_or_404(RegistroAuditoria.objects.select_related('usuario'), pk=pk)
    return render(request, 'facturacion/auditoria/auditoria_detalle.html', {
        'registro': registro,
    })


@login_required(login_url='login')
@user_passes_test(_es_superusuario, login_url='dashboard')
def auditoria_export_excel(request):
    """Exporta los registros de auditoría filtrados a Excel."""
    registros = RegistroAuditoria.objects.select_related('usuario').all().order_by('-fecha')
    registros, _ = _aplicar_filtros(request, registros)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Auditoría'

    encabezados = ['Fecha', 'Usuario', 'Acción', 'Modelo', 'Objeto ID', 'Objeto', 'IP', 'Navegador']
    ws.append(encabezados)

    relleno = PatternFill(start_color='16232C', end_color='16232C', fill_type='solid')
    fuente = Font(color='FFFFFF', bold=True)
    for col in range(1, len(encabezados) + 1):
        celda = ws.cell(row=1, column=col)
        celda.fill = relleno
        celda.font = fuente
        celda.alignment = Alignment(horizontal='center')

    for r in registros:
        ws.append([
            r.fecha.strftime('%d/%m/%Y %H:%M:%S'),
            r.usuario.username if r.usuario else 'Sistema',
            r.get_accion_display(),
            r.modelo,
            r.objeto_id,
            r.objeto_repr,
            r.ip or '',
            (r.user_agent or '')[:80],
        ])

    anchos = [20, 15, 15, 18, 10, 40, 15, 50]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="auditoria.xlsx"'
    wb.save(response)
    return response